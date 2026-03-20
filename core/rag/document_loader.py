"""
Sentinel AI — Document Loader

Loads documents from various formats (PDF, TXT, MD, DOCX) and
splits them into chunks suitable for embedding and retrieval.

Supports:
  - .txt, .md, .rst  — plain text
  - .pdf             — PDF via PyMuPDF (fitz) or pdfplumber fallback
  - .docx            — DOCX via python-docx
  - .json            — JSON array or object of text fields
  - Directories      — recursive loading with glob patterns

Chunking strategy: sentence-aware sliding window with configurable
chunk_size (tokens/chars) and overlap.
"""
import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Union

logger = logging.getLogger(__name__)

# Default chunking parameters
DEFAULT_CHUNK_SIZE = 512   # characters
DEFAULT_CHUNK_OVERLAP = 64  # characters


# ─────────────────────────────────────────────
# Document dataclass
# ─────────────────────────────────────────────

class Document:
    """
    A single document chunk with text and metadata.

    Attributes:
        text: The text content of this chunk.
        metadata: Source file, chunk index, page number, etc.
        doc_id: Unique identifier for this chunk.
    """

    def __init__(
        self,
        text: str,
        metadata: Optional[Dict[str, Any]] = None,
        doc_id: Optional[str] = None,
    ):
        import uuid
        self.text = text.strip()
        self.metadata = metadata or {}
        self.doc_id = doc_id or str(uuid.uuid4())

    def __repr__(self) -> str:
        preview = self.text[:60].replace("\n", " ")
        return f"Document(id={self.doc_id[:8]}…, text='{preview}…', meta={self.metadata})"

    def __len__(self) -> int:
        return len(self.text)


# ─────────────────────────────────────────────
# Chunking
# ─────────────────────────────────────────────

def chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
    source: str = "",
) -> List[Document]:
    """
    Split text into overlapping chunks.

    Strategy:
    1. Split on paragraph boundaries (double newlines) first.
    2. If a paragraph is still larger than chunk_size, split by sentences.
    3. Assemble chunks with overlap.

    Args:
        text: Full document text.
        chunk_size: Target chunk size in characters.
        overlap: Overlap between consecutive chunks in characters.
        source: Source identifier for metadata.

    Returns:
        List of Document objects.
    """
    if not text or not text.strip():
        return []

    # Normalize whitespace
    text = re.sub(r"\n{3,}", "\n\n", text.strip())

    # Split into paragraphs
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    # Reassemble into chunks respecting chunk_size
    chunks: List[str] = []
    current = ""

    for para in paragraphs:
        if len(current) + len(para) + 2 <= chunk_size:
            current = (current + "\n\n" + para).strip()
        else:
            if current:
                chunks.append(current)
            # Handle oversized paragraphs
            if len(para) > chunk_size:
                # Split by sentence
                sentences = re.split(r"(?<=[.!?])\s+", para)
                current = ""
                for sent in sentences:
                    if len(current) + len(sent) + 1 <= chunk_size:
                        current = (current + " " + sent).strip()
                    else:
                        if current:
                            chunks.append(current)
                        current = sent[:chunk_size]  # hard cut for very long sentences
            else:
                current = para

    if current:
        chunks.append(current)

    # Add overlap between consecutive chunks
    result: List[Document] = []
    for i, chunk in enumerate(chunks):
        # Prepend tail of previous chunk
        if i > 0 and overlap > 0:
            prev_tail = chunks[i - 1][-overlap:]
            chunk = prev_tail + " " + chunk

        result.append(Document(
            text=chunk,
            metadata={
                "source": source,
                "chunk_index": i,
                "total_chunks": len(chunks),
                "chunk_size": len(chunk),
            },
        ))

    return result


# ─────────────────────────────────────────────
# Loaders by file type
# ─────────────────────────────────────────────

def _load_text(path: Path) -> str:
    """Load plain text, markdown, RST, etc."""
    return path.read_text(encoding="utf-8", errors="replace")


def _load_pdf(path: Path) -> str:
    """Load PDF using PyMuPDF (preferred) or pdfplumber (fallback)."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        pages = [page.get_text() for page in doc]
        doc.close()
        return "\n\n".join(pages)
    except ImportError:
        pass

    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n\n".join(pages)
    except ImportError:
        raise ImportError(
            "PDF support requires PyMuPDF or pdfplumber. "
            "Install with: pip install pymupdf OR pip install pdfplumber"
        )


def _load_docx(path: Path) -> str:
    """Load DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        raise ImportError(
            "DOCX support requires python-docx. "
            "Install with: pip install python-docx"
        )


def _load_json(path: Path) -> str:
    """Load JSON — extracts text values from dicts or array items."""
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        parts = []
        for item in data:
            if isinstance(item, dict):
                # Concatenate all string values
                parts.append(" ".join(str(v) for v in item.values() if isinstance(v, str)))
            elif isinstance(item, str):
                parts.append(item)
        return "\n\n".join(parts)
    elif isinstance(data, dict):
        return "\n\n".join(str(v) for v in data.values() if isinstance(v, str))
    return str(data)


# ─────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────

SUPPORTED_EXTENSIONS = {".txt", ".md", ".rst", ".pdf", ".docx", ".json"}


def load_document(
    path: Union[str, Path],
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> List[Document]:
    """
    Load a document from a file and split into chunks.

    Args:
        path: Path to the document.
        chunk_size: Target chunk size in characters.
        overlap: Overlap between chunks in characters.

    Returns:
        List of Document chunks.

    Raises:
        FileNotFoundError: If the file doesn't exist.
        ValueError: If the file extension is not supported.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    logger.info(f"Loading document: {path.name}")

    if ext in {".txt", ".md", ".rst"}:
        text = _load_text(path)
    elif ext == ".pdf":
        text = _load_pdf(path)
    elif ext == ".docx":
        text = _load_docx(path)
    elif ext == ".json":
        text = _load_json(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")

    chunks = chunk_text(text, chunk_size=chunk_size, overlap=overlap, source=str(path))
    logger.info(f"Split '{path.name}' into {len(chunks)} chunks")
    return chunks


def load_directory(
    directory: Union[str, Path],
    glob: str = "**/*",
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
    exclude_patterns: Optional[List[str]] = None,
) -> List[Document]:
    """
    Recursively load all supported documents from a directory.

    Args:
        directory: Path to the directory.
        glob: Glob pattern for file matching (default: all files recursively).
        chunk_size: Target chunk size in characters.
        overlap: Overlap between chunks in characters.
        exclude_patterns: File name patterns to skip (e.g., ["secret*"]).

    Returns:
        List of all Document chunks from all files.
    """
    directory = Path(directory)
    if not directory.is_dir():
        raise NotADirectoryError(f"Not a directory: {directory}")

    exclude = set(exclude_patterns or [])
    all_docs: List[Document] = []
    total_files = 0

    for file_path in sorted(directory.glob(glob)):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        if any(file_path.match(pat) for pat in exclude):
            continue

        try:
            docs = load_document(file_path, chunk_size=chunk_size, overlap=overlap)
            all_docs.extend(docs)
            total_files += 1
        except Exception as e:
            logger.warning(f"Skipping {file_path.name}: {e}")

    logger.info(
        f"Loaded {len(all_docs)} chunks from {total_files} files in '{directory}'"
    )
    return all_docs
